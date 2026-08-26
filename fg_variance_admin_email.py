"""
scripts/sap_recon_email.py

SAP Recon Email Reporter — LOCAL TEST VERSION
=============================================
All credentials hardcoded. Run directly in VS Code:
    python scripts/sap_recon_email.py

Once verified, keep this file and trigger via GitHub Actions on a schedule.
"""

import base64
import logging
import sys
import traceback
from datetime import datetime
from io import BytesIO
from typing import Dict, List
from urllib.parse import quote

import msal
import psycopg2
import psycopg2.extras
import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ============================================================
# LOGGING
# ============================================================
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)-8s | %(message)s",
    handlers=[logging.StreamHandler()],
)
logger = logging.getLogger(__name__)

# ============================================================
# HARDCODED CONFIG
# ============================================================

# ── Database ─────────────────────────────────────────────────
DATABASE_URL = "postgres://udaeok46kp345o:p16a757221737effd08a392567cdd32730fde4632adf21803a2bb5127e5d2855e@c4717h2lt37nua.cluster-czrs8kj4isg7.us-east-1.rds.amazonaws.com:5432/d1faihjj66t25t"
# Get it by running: heroku config:get DATABASE_URL --app <your-app-name>
# It looks like: postgresql://user:pass@host:5432/dbname

# ── SAP OData ────────────────────────────────────────────────
SAP_BASE_URL  = "https://my409486-api.s4hana.cloud.sap/sap/opu/odata/sap/API_MATERIAL_STOCK_SRV"
SAP_USERNAME  = "POWERBI_INTEGRATION"
SAP_PASSWORD  = "d>xEX$dgfFjjw2=2<kbNkH-gPyu-+XnVJF%xJr)v"
SAP_PAGE_SIZE = 1000
SAP_TIMEOUT   = 300
SAP_SELECT    = "Material,Plant,MatlWrhsStkQtyInMatlBaseUnit"
SAP_VALID_PLANTS  = frozenset({"1000", "2000"})
SAP_INDUS_LETTERS = frozenset("ACI")

# ── SharePoint / Graph (OOR customer codes) ──────────────────
SP_TENANT_ID     = "ed97e9bb-e119-4bd4-ab00-307d64bdf908"
SP_CLIENT_ID     = "a2239be7-12cd-442d-983f-ea7e316ef767"
SP_CLIENT_SECRET = "sho8Q~sanlYEX3SKXjYJb0NSXL.gPRpdi63-tcmp"
SP_SITE_ID       = "201c767f-23cc-4eff-b76c-47a275706531,094b21f9-6ed4-47c3-a751-3b217a18b98e"
SP_OOR_FILENAME  = "Indus Open Order Report New.xlsx"

# ── Mail (Graph API) ─────────────────────────────────────────
CLIENT_ID = "785b5fc9-1624-458d-90cb-a140dff579c7"
CLIENT_SECRET = "Zcw8Q~YSEqXPypF4Gn_mNVpp2MVPRMM.95hjVdbP"
AUTHORITY = "https://login.microsoftonline.com/ed97e9bb-e119-4bd4-ab00-307d64bdf908"

SENDER_EMAIL = "siddhesh.mandhana@indusair.ae"
TO_EMAILS = [
    "abhir.bhomavat@indusair.ae",
    "aanand.sr@indusair.ae",
    "brijendu.pandey@indusair.ae",
    "shabbir.kachwala@indusair.ae",
]
CC_EMAILS = [
    "meet.shah@indusair.ae",
    "prodndbx3@indusair.ae",
    "ppc2@indusair.ae",
    "sachin.nanekar@indusair.ae",
    "prodndbx11@indusair.ae",
    "Shivaji.Korde@indusair.ae",
]

# ── Change these to just yourself while testing ──────────────
# TO_EMAILS   = ["siddhesh.mandhana@indusair.ae"]
# CC_EMAILS   = []

from datetime import timezone, timedelta

_DUBAI_TZ = timezone(timedelta(hours=4))
_NOW_DUBAI = datetime.now(_DUBAI_TZ)

EMAIL_SUBJECT       = f"Indus Warehouse-SAP FG Variance [{_NOW_DUBAI.strftime('%d-%b-%Y')}]"
ATTACHMENT_FILENAME = f"Indus Warehouse-SAP FG Variance {_NOW_DUBAI.strftime('%d-%b-%Y %H.%M')} GST.xlsx"


# ============================================================
# STEP 1 — POSTGRES: scanner net qty per indus_part
# ============================================================

def fetch_scanner_data() -> Dict[str, dict]:
    logger.info("🗄️  Connecting to Postgres...")
    sql = """
        SELECT
            UPPER(COALESCE(NULLIF(indus_part, ''), 'Unknown')) AS indus_part,
            COALESCE(NULLIF(MAX(customer_part), ''), '')       AS customer_part,
            SUM(CASE WHEN movement_type = 'STOCK_IN'
                     THEN COALESCE(pallet_qty, 0) ELSE 0 END)
            - SUM(CASE WHEN movement_type = 'STOCK_OUT'
                     THEN COALESCE(pallet_qty, 0) ELSE 0 END)  AS net_qty
        FROM movements
        WHERE movement_type IN ('STOCK_IN', 'STOCK_OUT')
        GROUP BY UPPER(COALESCE(NULLIF(indus_part, ''), 'Unknown'))
        ORDER BY UPPER(COALESCE(NULLIF(indus_part, ''), 'Unknown'))
    """
    conn = psycopg2.connect(DATABASE_URL, sslmode="require")
    try:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(sql)
            rows = cur.fetchall()
        result = {row["indus_part"]: dict(row) for row in rows}
        logger.info("✅  Scanner data: %d materials", len(result))
        return result
    finally:
        conn.close()


# ============================================================
# STEP 2 — SAP OData: plant stock for plants 1000 + 2000
# ============================================================

def fetch_sap_stock() -> Dict[str, dict]:
    from requests.auth import HTTPBasicAuth

    logger.info("🔗  Fetching SAP stock (this takes 2-5 min, be patient)...")
    endpoint = f"{SAP_BASE_URL}/A_MatlStkInAcctMod"
    auth     = HTTPBasicAuth(SAP_USERNAME, SAP_PASSWORD)
    headers  = {"Accept": "application/json"}
    all_records, skip, page = [], 0, 0

    while True:
        page += 1
        params = {
            "$format": "json",
            "$select": SAP_SELECT,
            "$top":    str(SAP_PAGE_SIZE),
            "$skip":   str(skip),
        }
        logger.info("  Page %d | skip=%d | total so far: %d", page, skip, len(all_records))
        resp = requests.get(endpoint, params=params, headers=headers,
                            auth=auth, timeout=SAP_TIMEOUT)
        resp.raise_for_status()
        results = resp.json().get("d", {}).get("results", [])
        all_records.extend(results)
        if len(results) < SAP_PAGE_SIZE:
            logger.info("  Last page reached — %d total raw records", len(all_records))
            break
        skip += SAP_PAGE_SIZE

    # Filter + aggregate (same logic as sap_stock_client.py)
    agg: Dict[str, dict] = {}
    for rec in all_records:
        plant = (rec.get("Plant") or "").strip()
        if plant not in SAP_VALID_PLANTS:
            continue
        mat = (rec.get("Material") or "").strip().upper()
        if not mat or not any(l in mat for l in SAP_INDUS_LETTERS):
            continue
        try:
            qty = float(rec.get("MatlWrhsStkQtyInMatlBaseUnit") or 0)
        except (TypeError, ValueError):
            qty = 0.0
        if mat not in agg:
            agg[mat] = {"material": mat, "sap_qty_1000": 0.0, "sap_qty_2000": 0.0}
        if plant == "1000":
            agg[mat]["sap_qty_1000"] += qty
        elif plant == "2000":
            agg[mat]["sap_qty_2000"] += qty

    logger.info("✅  SAP stock: %d unique Indus materials after filter", len(agg))
    return agg


# ============================================================
# STEP 3 — SharePoint: OOR customer codes (col H → col R)
# ============================================================

def _get_sp_token() -> str:
    app = msal.ConfidentialClientApplication(
        SP_CLIENT_ID,
        authority=f"https://login.microsoftonline.com/{SP_TENANT_ID}",
        client_credential=SP_CLIENT_SECRET,
    )
    result = app.acquire_token_for_client(scopes=["https://graph.microsoft.com/.default"])
    if "access_token" not in result:
        raise Exception(result.get("error_description", result.get("error")))
    return result["access_token"]


def fetch_oor_customer_codes() -> Dict[str, str]:
    logger.info("📋  Fetching OOR customer codes from SharePoint...")
    token   = _get_sp_token()
    headers = {"Authorization": f"Bearer {token}"}

    # Find the workbook
    search_url = (
        f"https://graph.microsoft.com/v1.0/sites/{SP_SITE_ID}"
        f"/drive/root/search(q='{SP_OOR_FILENAME}')"
    )
    items = requests.get(search_url, headers=headers, timeout=30).json().get("value", [])
    item  = next(
        (i for i in items if i.get("name", "").lower() == SP_OOR_FILENAME.lower()), None
    )
    if not item:
        logger.warning("⚠️  OOR workbook not found — customer column will be blank")
        return {}

    item_id = item["id"]

    # Get first/matching worksheet
    sheets = requests.get(
        f"https://graph.microsoft.com/v1.0/sites/{SP_SITE_ID}"
        f"/drive/items/{item_id}/workbook/worksheets",
        headers=headers, timeout=30,
    ).json().get("value", [])
    if not sheets:
        return {}
    sheet_name = sheets[0]["name"]
    for s in sheets:
        if "open order" in s.get("name", "").lower():
            sheet_name = s["name"]
            break

    # Read used range
    all_rows = requests.get(
        f"https://graph.microsoft.com/v1.0/sites/{SP_SITE_ID}"
        f"/drive/items/{item_id}/workbook/worksheets('{quote(sheet_name)}')/usedRange",
        headers=headers, timeout=120,
    ).json().get("values", [])

    result: Dict[str, set] = {}
    for row in all_rows[1:]:
        if len(row) <= 17:
            continue
        part = str(row[7]).strip().upper()
        code = str(row[17]).strip().upper()
        if not part or part in ("", "NONE", "N/A", "#N/A"):
            continue
        if not code or code in ("", "NONE", "N/A", "#N/A"):
            continue
        result.setdefault(part, set()).add(code)

    merged = {k: ", ".join(sorted(v)) for k, v in result.items()}
    logger.info("✅  OOR: %d parts mapped", len(merged))
    return merged


# ============================================================
# STEP 4 — Build recon rows (full outer join)
# ============================================================

def build_recon_rows(
    scanner_map: Dict[str, dict],
    sap_map:     Dict[str, dict],
    oor_map:     Dict[str, str],
) -> List[dict]:
    all_materials = sorted(set(sap_map.keys()) | set(scanner_map.keys()))
    rows = []
    for mat in all_materials:
        sap  = sap_map.get(mat, {})
        scan = scanner_map.get(mat, {})
        sq   = int(scan.get("net_qty", 0) or 0)
        s1   = round(float(sap.get("sap_qty_1000", 0) or 0))
        s2   = round(float(sap.get("sap_qty_2000", 0) or 0))
        rows.append({
            "indus_part":       mat,
            "customer":         oor_map.get(mat, ""),
            "customer_part":    scan.get("customer_part", ""),
            "scanner_qty":      sq,
            "sap_qty_1000":     s1,
            "sap_qty_2000":     s2,
            "variance_vs_1000": sq - s1,
        })
    logger.info("✅  %d recon rows built (%d discrepant)",
                len(rows), sum(1 for r in rows if r["variance_vs_1000"] != 0))
    return rows


# ============================================================
# STEP 5 — Build Excel
# ============================================================

def build_excel(rows: List[dict]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "SAP Recon"

    HDR_FILL   = PatternFill("solid", fgColor="1E3A5F")
    HDR_FONT   = Font(bold=True, color="FFFFFF", size=10, name="Calibri")
    HDR_ALIGN  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    DATA_FONT  = Font(size=10, name="Calibri")
    DATA_ALIGN = Alignment(vertical="center")
    _thin      = Side(style="thin", color="CCCCCC")
    BORDER     = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
    NEG_FILL   = PatternFill("solid", fgColor="2D0A0A")
    NEG_FONT   = Font(size=10, name="Calibri", color="FCA5A5")

    HEADERS  = ["indus_part", "customer", "customer_part",
                "scanner_qty", "sap_qty_1000", "sap_qty_2000", "variance_vs_1000"]
    INT_COLS = frozenset({"scanner_qty", "sap_qty_1000", "sap_qty_2000", "variance_vs_1000"})

    ws.append(HEADERS)
    ws.row_dimensions[1].height = 22
    for cell in ws[1]:
        cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.alignment = HDR_ALIGN; cell.border = BORDER
    for j in range(1, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 22

    for i, row in enumerate(rows, start=2):
        ws.row_dimensions[i].height = 18
        is_neg = int(row.get("variance_vs_1000", 0)) < 0
        for j, h in enumerate(HEADERS, start=1):
            val      = row.get(h)
            cell_val = int(val) if h in INT_COLS else (str(val) if val is not None else "")
            cell     = ws.cell(row=i, column=j, value=cell_val)
            cell.font      = NEG_FONT if is_neg else DATA_FONT
            cell.fill      = NEG_FILL if is_neg else PatternFill()
            cell.alignment = DATA_ALIGN
            cell.border    = BORDER

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    logger.info("✅  Excel built: %d rows", len(rows))
    return out


# ============================================================
# STEP 6 — Build HTML email body
# ============================================================

def build_html_body(rows: List[dict]) -> str:
    run_date  = datetime.now().strftime("%d-%b-%Y %H:%M")
    disc_rows = [r for r in rows if r["variance_vs_1000"] != 0]
    total_var = sum(r["variance_vs_1000"] for r in disc_rows)

    def fmt(v):
        try:
            fv = float(v)
            return f"{int(fv):,}" if fv == int(fv) else f"{fv:,.2f}"
        except (ValueError, TypeError):
            return str(v)

    def fmt_var(v):
        try:
            fv = float(v)
            return f"{int(fv):+,}" if fv == int(fv) else f"{fv:+,.2f}"
        except (ValueError, TypeError):
            return str(v)

    total_color = "#C0392B" if total_var < 0 else ("#E67E22" if total_var > 0 else "#27AE60")

    banner = f"""
    <div style="background:#1E3A5F;color:#fff;border-radius:6px;padding:12px 16px;margin-bottom:16px;">
        <strong>Total Variance (Plant 1000):</strong>
        <span style="color:{total_color};font-size:16px;font-weight:bold;margin-left:10px;">
            {fmt_var(total_var)}
        </span>
        &nbsp;|&nbsp; <strong>{len(disc_rows)}</strong> discrepant materials
        &nbsp;|&nbsp; Run: {run_date}
    </div>"""

    if not disc_rows:
        table_html = """<div style="text-align:center;padding:20px;color:#27AE60;font-size:16px;">
            ✅ No variance exceptions — all materials are balanced.</div>"""
    else:
        col_cfg = [
            ("indus_part",       "120px", "#1F3864"),
            ("customer",         "110px", "#1F3864"),
            ("customer_part",    "110px", "#1F3864"),
            ("scanner_qty",       "90px", "#1F3864"),
            ("sap_qty_1000",      "90px", "#1F3864"),
            ("sap_qty_2000",      "90px", "#1F3864"),
            ("variance_vs_1000", "130px", "#C0392B"),
        ]
        header_row = "".join(
            f'<th style="background:{bg};color:#fff;padding:6px 8px;'
            f'width:{w};text-align:center;border:1px solid #444;">{h}</th>'
            for h, w, bg in col_cfg
        )
        data_rows_html = ""
        for r in disc_rows:
            var = r["variance_vs_1000"]
            row_bg    = "#2d0a0a" if var < 0 else ("#3d2800" if var > 0 else "#ffffff")
            row_color = "#fca5a5" if var < 0 else ("#fde68a" if var > 0 else "#1f1f1f")
            cells = "".join(
                f'<td style="padding:5px 8px;border:1px solid #444;color:{row_color};'
                f'text-align:{"left" if h == "indus_part" else "center"};">'
                f'{fmt_var(r.get(h,"")) if h == "variance_vs_1000" else fmt(r.get(h,"")) if h in {"scanner_qty","sap_qty_1000","sap_qty_2000"} else str(r.get(h,""))}'
                f'</td>'
                for h, _, _ in col_cfg
            )
            data_rows_html += f'<tr style="background:{row_bg};">{cells}</tr>'

        table_html = f"""
        <table style="border-collapse:collapse;width:100%;font-size:12px;font-family:Calibri,sans-serif;">
            <thead><tr>{header_row}</tr></thead>
            <tbody>{data_rows_html}</tbody>
        </table>"""

    return f"""<html><body style="font-family:Calibri,Arial,sans-serif;font-size:13px;background:#f5f5f5;padding:20px;">
    <div style="max-width:900px;margin:auto;background:#fff;padding:24px;border-radius:8px;">
        <h2 style="color:#1E3A5F;margin-bottom:4px;">Indus SAP Reconciliation Report</h2>
        <p style="color:#555;margin-top:0;">Scanner Qty vs SAP Plant Stock (MB52) — discrepancies only shown below</p>
        {banner}
        <p style="margin-top:20px;font-size:13px;color:#1f1f1f;">
            <a href="mailto:aanand.sr@indusair.ae" style="color:#1E3A5F;font-weight:bold;text-decoration:none;">@Aanand SR</a>
            &nbsp;please advise on how to tackle these variances promptly.
        </p>
        {table_html}
        <br>
        <p style="color:#777;font-size:11px;margin-top:12px;">Auto-generated by Indus Warehouse Tracker. Full data in attached Excel.</p>
        <br>
        <p style="font-size:13px;color:#1f1f1f;line-height:1.8;">
            Best Regards,<br>
            <strong>Siddhesh Mandhana</strong><br>
            SAP S/4HANA Consultant &amp; Process Automation Lead<br>
            Indus International FZC
        </p>
    </div>
    </body></html>"""


# ============================================================
# STEP 7 — Send via Microsoft Graph API
# ============================================================

def _get_mail_token() -> str:
    app = msal.ConfidentialClientApplication(
        CLIENT_ID,
        authority=AUTHORITY,
        client_credential=CLIENT_SECRET,
    )
    result = app.acquire_token_for_client(scopes=["https://graph.microsoft.com/.default"])
    if "access_token" not in result:
        raise Exception(result.get("error_description", result.get("error")))
    return result["access_token"]


def send_email(html_body: str, excel_bytes: BytesIO):
    logger.info("📧  Sending email via Graph API...")
    token      = _get_mail_token()
    excel_b64  = base64.b64encode(excel_bytes.read()).decode("utf-8")

    payload = {
        "message": {
            "subject": EMAIL_SUBJECT,
            "body": {"contentType": "HTML", "content": html_body},
            "toRecipients": [{"emailAddress": {"address": e}} for e in TO_EMAILS],
            "ccRecipients": [{"emailAddress": {"address": e}} for e in CC_EMAILS],
            "attachments": [{
                "@odata.type":  "#microsoft.graph.fileAttachment",
                "name":         ATTACHMENT_FILENAME,
                "contentType":  "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "contentBytes": excel_b64,
            }],
        }
    }

    resp = requests.post(
        f"https://graph.microsoft.com/v1.0/users/{SENDER_EMAIL}/sendMail",
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        json=payload,
        timeout=60,
    )
    if resp.status_code == 202:
        logger.info("✅  Email sent to %s", TO_EMAILS)
    else:
        raise Exception(f"Graph API {resp.status_code}: {resp.text[:400]}")


# ============================================================
# MAIN
# ============================================================

def main():
    logger.info("=" * 60)
    logger.info("SAP Recon Email — %s", datetime.now().strftime("%d-%b-%Y %H:%M"))
    logger.info("=" * 60)

    try:
        scanner_map = fetch_scanner_data()
    except Exception as e:
        logger.error("❌  Postgres failed: %s", e); traceback.print_exc(); sys.exit(1)

    try:
        sap_map = fetch_sap_stock()
    except Exception as e:
        logger.error("❌  SAP fetch failed: %s", e); traceback.print_exc(); sys.exit(1)

    try:
        oor_map = fetch_oor_customer_codes()
    except Exception as e:
        logger.warning("⚠️  OOR skipped (%s) — customer column will be blank", e)
        oor_map = {}

    rows      = build_recon_rows(scanner_map, sap_map, oor_map)
    excel_out = build_excel(rows)
    html_body = build_html_body(rows)

    try:
        send_email(html_body, excel_out)
    except Exception as e:
        logger.error("❌  Email failed: %s", e); traceback.print_exc(); sys.exit(1)

    logger.info("🎉  Done.")


if __name__ == "__main__":
    main()
