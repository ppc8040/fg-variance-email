"""
INDUS PHYSICAL FG vs SAP — VARIANCE EMAIL REPORTER
====================================================
Reads "Physical FG Variance Report" worksheet from the SharePoint file
"Indus Ultimate Stock Recon.xlsx", filters rows where column J
(Final Variance Physical vs SAP) is non-zero, sends a formatted
HTML email via Microsoft Graph API, and attaches the same variance
report as a downloadable Excel file.

Columns included in report:
  A  – Indus Part
  B  – Customer
  E  – Physical Warehouse Total
  H  – Production FG Total
  I  – SAP Quantities (1000+2000)
  J  – Final Variance (Physical vs SAP)

The Excel attachment is formatted with borders, fills, alignment,
and a worksheet name derived from the email subject.

Author: Siddhesh Mandhana
"""

import base64
import re
import requests
import pandas as pd
import msal
import logging
import sys
import traceback
from pathlib import Path
from datetime import datetime
from urllib.parse import quote
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

# ============================================================
# SYSTEM SETUP
# ============================================================
if sys.platform == "win32":
    import os
    os.environ["PYTHONIOENCODING"] = "utf-8"

# ============================================================
# LOGGING
# ============================================================
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)-8s | %(message)s",
    handlers=[logging.StreamHandler()]
)
logger = logging.getLogger(__name__)

# ============================================================
# CONFIGURATION
# ============================================================
class SharePointAuthConfig:
    CLIENT_ID = "a2239be7-12cd-442d-983f-ea7e316ef767"
    CLIENT_SECRET = "sho8Q~sanlYEX3SKXjYJb0NSXL.gPRpdi63-tcmp"
    AUTHORITY = "https://login.microsoftonline.com/ed97e9bb-e119-4bd4-ab00-307d64bdf908"


class MailAuthConfig:
    CLIENT_ID = "785b5fc9-1624-458d-90cb-a140dff579c7"
    CLIENT_SECRET = "Zcw8Q~YSEqXPypF4Gn_mNVpp2MVPRMM.95hjVdbP"
    AUTHORITY = "https://login.microsoftonline.com/ed97e9bb-e119-4bd4-ab00-307d64bdf908"


class SharePointConfig:
    SITE_ID = "201c767f-23cc-4eff-b76c-47a275706531,094b21f9-6ed4-47c3-a751-3b217a18b98e"
    FILE_NAME = "Indus Ultimate Stock Recon.xlsx"
    WORKSHEET_NAME = "Physical FG Variance Report"
    HEADER_ROW_INDEX = 3
    TOTAL_VARIANCE_ROW_INDEX = 1
    TOTAL_VARIANCE_COL_INDEX = 9


class EmailConfig:
    SENDER_EMAIL = "releases@indusair.ae"
    TO_EMAILS = ["siddhesh.mandhana@indusair.ae"]
    CC_EMAILS = []
    SUBJECT = f"Indus Physical FG vs SAP — Variance Report [{datetime.now().strftime('%d-%b-%Y')}]"
    ATTACHMENT_FILE_NAME = f"Indus_Physical_FG_vs_SAP_Variance_Report_{datetime.now().strftime('%d-%b-%Y')}.xlsx"


# ============================================================
# AUTHENTICATION
# ============================================================
def get_sharepoint_token():
    logger.info("🔐 Acquiring SharePoint access token...")
    try:
        app = msal.ConfidentialClientApplication(
            SharePointAuthConfig.CLIENT_ID,
            authority=SharePointAuthConfig.AUTHORITY,
            client_credential=SharePointAuthConfig.CLIENT_SECRET
        )
        result = app.acquire_token_silent(
            scopes=["https://graph.microsoft.com/.default"],
            account=None
        )
        if not result:
            result = app.acquire_token_for_client(scopes=["https://graph.microsoft.com/.default"])
        if "access_token" in result:
            logger.info("✅ SharePoint token acquired")
            return result["access_token"]
        raise Exception(result.get("error_description", result.get("error")))
    except Exception as e:
        logger.error(f"❌ SharePoint token error: {e}")
        traceback.print_exc()
        raise


def get_mail_token():
    logger.info("🔐 Acquiring Mail send token...")
    try:
        app = msal.ConfidentialClientApplication(
            MailAuthConfig.CLIENT_ID,
            authority=MailAuthConfig.AUTHORITY,
            client_credential=MailAuthConfig.CLIENT_SECRET
        )
        result = app.acquire_token_for_client(scopes=["https://graph.microsoft.com/.default"])
        if "access_token" in result:
            logger.info("✅ Mail token acquired")
            return result["access_token"]
        raise Exception(result.get("error_description", result.get("error")))
    except Exception as e:
        logger.error(f"❌ Mail token error: {e}")
        traceback.print_exc()
        raise


# ============================================================
# SHAREPOINT READ
# ============================================================
def read_worksheet(sp_token):
    logger.info(f"📂 Reading worksheet: '{SharePointConfig.WORKSHEET_NAME}'")

    session_obj = requests.Session()
    api_headers = {
        "Authorization": f"Bearer {sp_token}",
        "Content-Type": "application/json"
    }
    site_id = SharePointConfig.SITE_ID
    file_name = quote(SharePointConfig.FILE_NAME)
    ws_name = quote(SharePointConfig.WORKSHEET_NAME)
    session_id = None

    try:
        session_url = (
            f"https://graph.microsoft.com/v1.0/sites/{site_id}"
            f"/drive/root:/{file_name}:/workbook/createSession"
        )
        session_resp = session_obj.post(
            session_url,
            headers=api_headers,
            json={"persistChanges": False},
            timeout=30
        )
        if session_resp.status_code not in (200, 201):
            raise Exception(f"Failed to create workbook session: {session_resp.status_code} — {session_resp.text[:300]}")

        session_id = session_resp.json().get("id")
        api_headers["workbook-session-id"] = session_id
        logger.info(f"✅ Workbook session created: {session_id[:30]}...")

        range_url = (
            f"https://graph.microsoft.com/v1.0/sites/{site_id}"
            f"/drive/root:/{file_name}:/workbook/worksheets('{ws_name}')/usedRange"
        )
        range_resp = session_obj.get(range_url, headers=api_headers, timeout=120)
        if range_resp.status_code != 200:
            raise Exception(f"Failed to read usedRange: {range_resp.status_code} — {range_resp.text[:300]}")

        values = range_resp.json().get("values", [])
        logger.info(f"✅ Raw rows returned from API: {len(values)}")
    except Exception as e:
        logger.error(f"❌ Worksheet read error: {e}")
        traceback.print_exc()
        raise
    finally:
        if session_id:
            try:
                close_url = (
                    f"https://graph.microsoft.com/v1.0/sites/{site_id}"
                    f"/drive/root:/{file_name}:/workbook/closeSession"
                )
                session_obj.post(close_url, headers=api_headers, timeout=15)
                logger.info("✅ Workbook session closed")
            except Exception as ce:
                logger.warning(f"⚠️ Could not close session: {ce}")

    total_variance = 0
    try:
        total_variance_raw = values[SharePointConfig.TOTAL_VARIANCE_ROW_INDEX][SharePointConfig.TOTAL_VARIANCE_COL_INDEX]
        total_variance = float(total_variance_raw) if total_variance_raw not in ("", None) else 0
        logger.info(f"📊 Total variance (J2): {total_variance}")
    except (IndexError, ValueError, TypeError) as e:
        logger.warning(f"⚠️ Could not read total variance cell: {e}. Defaulting to 0.")

    header_idx = SharePointConfig.HEADER_ROW_INDEX
    if len(values) <= header_idx:
        raise Exception(f"Not enough rows returned. Expected headers at index {header_idx}, but only {len(values)} row(s) found.")

    headers = values[header_idx]
    data_rows = values[header_idx + 1:]
    logger.info(f"📋 Headers detected: {headers}")
    logger.info(f"📋 Data rows: {len(data_rows)}")

    if not data_rows:
        raise Exception("No data rows found after header row.")

    df = pd.DataFrame(data_rows, columns=headers)
    logger.info(f"✅ DataFrame built: {df.shape[0]} rows × {df.shape[1]} columns")
    return df, total_variance


# ============================================================
# FILTER DATA
# ============================================================
def filter_variance_rows(df):
    logger.info("🔍 Selecting columns and filtering variance rows...")
    total_cols = df.shape[1]
    logger.info(f"   Total columns available: {total_cols}")

    if total_cols < 10:
        raise Exception(
            f"Expected at least 10 columns (A–J), but got {total_cols}. "
            f"Check HEADER_ROW_INDEX (currently {SharePointConfig.HEADER_ROW_INDEX})."
        )

    df_sel = df.iloc[:, [0, 1, 4, 7, 8, 9]].copy()
    df_sel.columns = [
        "Indus Part",
        "Customer",
        "Physical Warehouse Total",
        "Production FG Total",
        "SAP Quantities (1000+2000)",
        "Final Variance (Physical vs SAP)"
    ]

    before = len(df_sel)
    df_sel = df_sel[
        ~(
            df_sel["Indus Part"].astype(str).str.strip().isin(["", "nan"])
            & df_sel["Customer"].astype(str).str.strip().isin(["", "nan"])
        )
    ].copy()
    logger.info(f"   Dropped {before - len(df_sel)} blank rows. Remaining: {len(df_sel)}")

    numeric_cols = [
        "Physical Warehouse Total",
        "Production FG Total",
        "SAP Quantities (1000+2000)",
        "Final Variance (Physical vs SAP)"
    ]
    for col in numeric_cols:
        df_sel[col] = pd.to_numeric(df_sel[col], errors="coerce")

    df_sel["Final Variance (Physical vs SAP)"] = df_sel["Final Variance (Physical vs SAP)"].fillna(0)
    result = df_sel[df_sel["Final Variance (Physical vs SAP)"] != 0].copy()

    for col in ["Physical Warehouse Total", "Production FG Total", "SAP Quantities (1000+2000)"]:
        result[col] = result[col].fillna(0)

    logger.info(f"✅ Variance rows found: {len(result)} out of {len(df_sel)} total")
    if not result.empty:
        logger.info(f"   Sample:\n{result.head(5).to_string()}")
    return result


# ============================================================
# EMAIL HTML
# ============================================================
def build_html_body(variance_df, total_variance):
    run_date = datetime.now().strftime("%d-%b-%Y %H:%M")

    def variance_row_colour(val):
        try:
            v = float(val)
        except (ValueError, TypeError):
            return "#FFFFFF"
        if v < 0:
            return "#FFD7D7"
        if v > 0:
            return "#FFF3CD"
        return "#FFFFFF"

    def variance_cell_colour(val):
        try:
            v = float(val)
        except (ValueError, TypeError):
            return "#FFFFFF"
        if v < 0:
            return "#FF6B6B"
        if v > 0:
            return "#FFA500"
        return "#FFFFFF"

    def fmt(v):
        try:
            fv = float(v)
            return f"{int(fv):,}" if fv == int(fv) else f"{fv:,.2f}"
        except (ValueError, TypeError):
            return str(v)

    def fmt_var(v):
        try:
            fv = float(v)
            return f"{int(fv):+,}" if fv == int(fv) else f"{fv:+,.2f}"
        except (ValueError, TypeError):
            return str(v)

    total_color = "#C0392B" if total_variance < 0 else ("#E67E22" if total_variance > 0 else "#27AE60")
    total_label = fmt_var(total_variance)

    summary_banner = f"""
    <table width="100%" cellpadding="0" cellspacing="0" style="margin-bottom:22px;">
      <tr>
        <td style="background:#1F3864;color:#FFFFFF;padding:14px 20px;border-radius:6px 6px 0 0;
                   font-family:Calibri,Arial,sans-serif;font-size:18px;font-weight:bold;">
          📊 Physical FG vs SAP — Variance Summary
        </td>
      </tr>
      <tr>
        <td style="background:#F4F6FB;padding:14px 20px;border:1px solid #CBD5E0;border-top:none;
                   border-radius:0 0 6px 6px;font-family:Calibri,Arial,sans-serif;font-size:18px;line-height:1.6;">
          <strong>Report Date:</strong> {run_date}&nbsp;&nbsp;&nbsp;
          <strong>Variance Rows:</strong> {len(variance_df)}&nbsp;&nbsp;&nbsp;
          <strong>Grand Total Variance:</strong>
          <span style="color:{total_color};font-weight:bold;font-size:18px;">{total_label}</span>
          &nbsp;
          <span style="color:#888;font-size:18px;">(Negative = excess stock in SAP)</span>
        </td>
      </tr>
    </table>
    """

    if variance_df.empty:
        table_section = """
        <p style="font-family:Calibri,Arial,sans-serif;font-size:18px;color:#27AE60;font-weight:bold;">
          ✅ No variance exceptions found. All records are currently balanced.
        </p>
        """
    else:
        col_cfg = [
            ("Indus Part", "130px", "#1F3864", "#FFFFFF"),
            ("Customer", "110px", "#1F3864", "#FFFFFF"),
            ("Physical Warehouse Total", "150px", "#1F3864", "#FFFFFF"),
            ("Production FG Total", "145px", "#1F3864", "#FFFFFF"),
            ("SAP Quantities (1000+2000)", "155px", "#1F3864", "#FFFFFF"),
            ("Final Variance (Physical vs SAP)", "170px", "#C0392B", "#FFFFFF"),
        ]

        header_row = "".join(
            f'<th style="background:{bg};color:{fg};padding:10px 12px;font-family:Calibri,Arial,sans-serif;'
            f'font-size:14px;border:1px solid #A0AEC0;white-space:nowrap;width:{w};">{col}</th>'
            for col, w, bg, fg in col_cfg
        )

        data_rows_html = ""
        for _, row in variance_df.iterrows():
            var_val = row["Final Variance (Physical vs SAP)"]
            row_bg = variance_row_colour(var_val)
            var_bg = variance_cell_colour(var_val)
            cells = [
                f'<td style="background:{row_bg};padding:9px 12px;font-family:Calibri,Arial,sans-serif;font-size:16px;border:1px solid #CBD5E0;">{str(row["Indus Part"]).strip()}</td>',
                f'<td style="background:{row_bg};padding:9px 12px;font-family:Calibri,Arial,sans-serif;font-size:16px;border:1px solid #CBD5E0;">{str(row["Customer"]).strip()}</td>',
                f'<td style="background:{row_bg};padding:9px 12px;font-family:Calibri,Arial,sans-serif;font-size:16px;border:1px solid #CBD5E0;text-align:right;">{fmt(row["Physical Warehouse Total"])}</td>',
                f'<td style="background:{row_bg};padding:9px 12px;font-family:Calibri,Arial,sans-serif;font-size:16px;border:1px solid #CBD5E0;text-align:right;">{fmt(row["Production FG Total"])}</td>',
                f'<td style="background:{row_bg};padding:9px 12px;font-family:Calibri,Arial,sans-serif;font-size:16px;border:1px solid #CBD5E0;text-align:right;">{fmt(row["SAP Quantities (1000+2000)"])}</td>',
                f'<td style="background:{var_bg};color:#FFFFFF;font-weight:bold;padding:9px 12px;font-family:Calibri,Arial,sans-serif;font-size:16px;border:1px solid #CBD5E0;text-align:right;">{fmt_var(var_val)}</td>',
            ]
            data_rows_html += f'<tr>{"".join(cells)}</tr>\n'

        table_section = f"""
        <table cellpadding="0" cellspacing="0" style="border-collapse:collapse;width:100%;margin-bottom:20px;">
          <thead><tr>{header_row}</tr></thead>
          <tbody>{data_rows_html}</tbody>
        </table>
        <p style="font-family:Calibri,Arial,sans-serif;font-size:18px;color:#888;line-height:1.5;">
          🔴 Red rows = SAP has excess (negative variance) &nbsp;|&nbsp;
          🟠 Amber rows = Physical stock exceeds SAP (positive variance)
        </p>
        """

    html = f"""
<!DOCTYPE html>
<html>
<head><meta charset="UTF-8"></head>
<body style="margin:0;padding:0;background:#F0F4F8;">
<table width="100%" cellpadding="0" cellspacing="0" style="background:#F0F4F8;padding:24px;">
  <tr>
    <td align="center">
      <table width="920" cellpadding="0" cellspacing="0" style="background:#FFFFFF;border-radius:8px;box-shadow:0 2px 12px rgba(0,0,0,0.08);overflow:hidden;">
        <tr>
          <td style="background:#1F3864;padding:24px 30px;">
            <p style="margin:0;font-family:Calibri,Arial,sans-serif;font-size:28px;font-weight:bold;color:#FFFFFF;letter-spacing:0.3px;">
              Indus Physical FG vs SAP — Variance Report
            </p>
            <p style="margin:6px 0 0;font-family:Calibri,Arial,sans-serif;font-size:18px;color:#A8C4E0;">
              Worksheet: Physical FG Variance Report &nbsp;|&nbsp; Generated: {run_date}
            </p>
          </td>
        </tr>
        <tr>
          <td style="padding:28px 30px;">
            <p style="font-family:Calibri,Arial,sans-serif;font-size:18px;color:#2D3748;margin:0 0 18px;line-height:1.7;">
              Hi Team,
            </p>
            <p style="font-family:Calibri,Arial,sans-serif;font-size:18px;color:#2D3748;margin:0 0 22px;line-height:1.7;">
              Please find below the <strong>Physical FG vs SAP Stock Variance Report</strong>
              for <strong>{datetime.now().strftime('%d-%b-%Y')}</strong>.
              Only items with a <strong>non-zero Final Variance</strong> are listed.
              The same report is also attached as a formatted Excel file for download and review.
              Please investigate and resolve any discrepancies at the earliest, and ensure that the <strong>Production team, Warehouse Managers, and Storekeepers</strong> continue to evaluate and physically count the stock on a <strong>daily basis</strong>.
            </p>
            {summary_banner}
            {table_section}
            <p style="font-family:Calibri,Arial,sans-serif;font-size:18px;color:#2D3748;margin:26px 0 4px;line-height:1.7;">
              Regards,<br>
              <strong>Sharon Patil</strong><br>
              Logistics<br>
              Indus International FZC, Sharjah
            </p>
          </td>
        </tr>
        <tr>
          <td style="background:#F4F6FB;padding:14px 30px;border-top:1px solid #E2E8F0;">
            <p style="margin:0;font-family:Calibri,Arial,sans-serif;font-size:15px;color:#A0AEC0;line-height:1.5;">
              This is an automated report generated from the Indus Ultimate Stock Recon SharePoint workbook. Do not reply to this email directly.
            </p>
          </td>
        </tr>
      </table>
    </td>
  </tr>
</table>
</body>
</html>
"""
    return html


# ============================================================
# EXCEL ATTACHMENT CREATION
# ============================================================
def safe_sheet_name(name):
    cleaned = re.sub(r'[\\/*?:\[\]]', ' ', name).strip()
    return cleaned[:31] if len(cleaned) > 31 else cleaned


def create_excel_attachment(variance_df, total_variance):
    logger.info("📎 Creating formatted Excel attachment...")

    output_dir = Path.home() / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    file_path = output_dir / EmailConfig.ATTACHMENT_FILE_NAME

    wb = Workbook()
    ws = wb.active
    ws.title = safe_sheet_name(EmailConfig.SUBJECT)

    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    sub_font = Font(name="Calibri", size=12, bold=True, color="1F1F1F")
    header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    cell_font = Font(name="Calibri", size=12, color="000000")
    white_bold_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    note_font = Font(name="Calibri", size=11, italic=True, color="666666")

    title_fill = PatternFill("solid", fgColor="1F3864")
    summary_fill = PatternFill("solid", fgColor="DCE6F1")
    header_fill = PatternFill("solid", fgColor="1F3864")
    variance_header_fill = PatternFill("solid", fgColor="C0392B")
    negative_fill = PatternFill("solid", fgColor="FDE9E7")
    positive_fill = PatternFill("solid", fgColor="FFF2CC")
    negative_var_fill = PatternFill("solid", fgColor="E74C3C")
    positive_var_fill = PatternFill("solid", fgColor="F39C12")
    total_negative_fill = PatternFill("solid", fgColor="F4CCCC")
    total_positive_fill = PatternFill("solid", fgColor="FCE5CD")
    total_zero_fill = PatternFill("solid", fgColor="D9EAD3")

    thin_side = Side(style="thin", color="7F8C8D")
    medium_side = Side(style="medium", color="5D6D7E")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    medium_border = Border(left=medium_side, right=medium_side, top=medium_side, bottom=medium_side)

    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    right = Alignment(horizontal="right", vertical="center")

    columns = [
        "Indus Part",
        "Customer",
        "Physical Warehouse Total",
        "Production FG Total",
        "SAP Quantities (1000+2000)",
        "Final Variance (Physical vs SAP)"
    ]

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["F"].width = 28

    ws.merge_cells("A1:F1")
    ws["A1"] = "Indus Physical FG vs SAP — Variance Report"
    ws["A1"].font = title_font
    ws["A1"].fill = title_fill
    ws["A1"].alignment = center
    ws["A1"].border = medium_border
    ws.row_dimensions[1].height = 24

    ws.merge_cells("A2:F2")
    ws["A2"] = f"Generated on {datetime.now().strftime('%d-%b-%Y %H:%M')}"
    ws["A2"].font = sub_font
    ws["A2"].fill = summary_fill
    ws["A2"].alignment = center
    ws["A2"].border = thin_border
    ws.row_dimensions[2].height = 20

    ws["A3"] = "Total Variance"
    ws["B3"] = total_variance
    ws["A3"].font = sub_font
    ws["B3"].font = sub_font
    ws["A3"].fill = summary_fill
    ws["B3"].fill = total_negative_fill if total_variance < 0 else total_positive_fill if total_variance > 0 else total_zero_fill
    ws["A3"].alignment = center
    ws["B3"].alignment = right
    ws["B3"].number_format = '+#,##0;-#,##0;0'
    ws["A3"].border = thin_border
    ws["B3"].border = thin_border

    header_row = 5
    for idx, col_name in enumerate(columns, start=1):
        c = ws.cell(row=header_row, column=idx, value=col_name)
        c.font = header_font
        c.fill = variance_header_fill if idx == 6 else header_fill
        c.alignment = center
        c.border = medium_border

    data_start_row = header_row + 1

    if variance_df.empty:
        ws.merge_cells(start_row=data_start_row, start_column=1, end_row=data_start_row, end_column=6)
        c = ws.cell(row=data_start_row, column=1, value="No non-zero variance rows found.")
        c.font = sub_font
        c.alignment = center
        c.border = thin_border
        ws.row_dimensions[data_start_row].height = 22
        last_row = data_start_row
    else:
        for r_idx, (_, row) in enumerate(variance_df.iterrows(), start=data_start_row):
            var_val = float(row["Final Variance (Physical vs SAP)"])
            base_fill = negative_fill if var_val < 0 else positive_fill
            var_fill = negative_var_fill if var_val < 0 else positive_var_fill

            values = [
                row["Indus Part"],
                row["Customer"],
                row["Physical Warehouse Total"],
                row["Production FG Total"],
                row["SAP Quantities (1000+2000)"],
                row["Final Variance (Physical vs SAP)"]
            ]

            for c_idx, value in enumerate(values, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                cell.font = white_bold_font if c_idx == 6 else cell_font
                cell.alignment = left if c_idx <= 2 else right
                cell.border = thin_border
                if c_idx == 6:
                    cell.fill = var_fill
                    cell.number_format = '+#,##0;-#,##0;0'
                else:
                    cell.fill = base_fill
                    if c_idx >= 3:
                        cell.number_format = '#,##0'
            ws.row_dimensions[r_idx].height = 20
        last_row = data_start_row + len(variance_df) - 1

    note_row = last_row + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)
    ws.cell(row=note_row, column=1, value="Red-shaded rows indicate SAP excess (negative variance). Amber-shaded rows indicate physical stock excess (positive variance).")
    ws.cell(row=note_row, column=1).font = note_font
    ws.cell(row=note_row, column=1).alignment = left

    ws.freeze_panes = "A6"
    if last_row >= data_start_row:
        ws.auto_filter.ref = f"A5:F{last_row}"

    wb.save(file_path)
    logger.info(f"✅ Excel attachment created: {file_path}")
    return file_path


# ============================================================
# EMAIL SEND
# ============================================================
def send_email(mail_token, html_body, attachment_path):
    logger.info("📧 Sending email via Microsoft Graph API...")

    to_recipients = [{"emailAddress": {"address": e}} for e in EmailConfig.TO_EMAILS]
    cc_recipients = [{"emailAddress": {"address": e}} for e in EmailConfig.CC_EMAILS]

    attachment_bytes = Path(attachment_path).read_bytes()
    attachment_b64 = base64.b64encode(attachment_bytes).decode("utf-8")

    payload = {
        "message": {
            "subject": EmailConfig.SUBJECT,
            "body": {
                "contentType": "HTML",
                "content": html_body
            },
            "toRecipients": to_recipients,
            "ccRecipients": cc_recipients,
            "attachments": [
                {
                    "@odata.type": "#microsoft.graph.fileAttachment",
                    "name": EmailConfig.ATTACHMENT_FILE_NAME,
                    "contentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    "contentBytes": attachment_b64
                }
            ]
        },
        "saveToSentItems": True
    }

    send_url = f"https://graph.microsoft.com/v1.0/users/{EmailConfig.SENDER_EMAIL}/sendMail"
    headers = {
        "Authorization": f"Bearer {mail_token}",
        "Content-Type": "application/json"
    }

    resp = requests.post(send_url, headers=headers, json=payload, timeout=60)
    if resp.status_code == 202:
        logger.info("✅ Email sent successfully with Excel attachment")
    else:
        raise Exception(f"Email send failed: {resp.status_code} — {resp.text[:400]}")


# ============================================================
# MAIN
# ============================================================
def main():
    logger.info("=" * 60)
    logger.info("  Indus Physical FG vs SAP — Variance Email Reporter")
    logger.info("=" * 60)

    try:
        sp_token = get_sharepoint_token()
        mail_token = get_mail_token()
        df, total_variance = read_worksheet(sp_token)
        variance_df = filter_variance_rows(df)
        html_body = build_html_body(variance_df, total_variance)
        attachment_path = create_excel_attachment(variance_df, total_variance)
        send_email(mail_token, html_body, attachment_path)
        logger.info("🎉 Done! Variance report email with Excel attachment dispatched.")
    except Exception as e:
        logger.error(f"💥 Fatal error: {e}")
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()
